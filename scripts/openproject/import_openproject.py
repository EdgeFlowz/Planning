from pathlib import Path
import os

import requests
import yaml
from openpyxl import load_workbook


REQUIRED_COLUMNS = {
    "Phase",
    "Type",
    "Subject",
    "Description",
    "Priority",
    "Estimated hours",
    "Owner",
    "Parent",
}


class OpenProjectClient:
    def __init__(self, base_url: str, api_token: str):
        self.base_url = base_url.rstrip("/")
        self.session = requests.Session()
        self.session.headers.update(
            {
                "Authorization": f"Bearer {api_token}",
                "Accept": "application/hal+json",
                "Content-Type": "application/json",
            }
        )

    def get_project(self, project_identifier: str) -> dict:
        response = self.session.get(
            f"{self.base_url}/api/v3/projects/{project_identifier}"
        )
        response.raise_for_status()
        return response.json()

    def get_work_packages(self, project_id: int) -> list[dict]:
        response = self.session.get(
            f"{self.base_url}/api/v3/projects/{project_id}/work_packages",
            params={"pageSize": 200},
        )
        response.raise_for_status()

        data = response.json()
        return data.get("_embedded", {}).get("elements", [])

    def create_work_package(
        self,
        project_id: int,
        work_type_id: int,
        subject: str,
        description: str | None,
        priority_id: int | None,
        estimated_hours: float | None,
        parent_id: int | None = None,
    ) -> dict:

        payload = {
            "subject": subject,
            "_links": {
                "type": {
                    "href": f"/api/v3/types/{work_type_id}"
                },
                "project": {
                    "href": f"/api/v3/projects/{project_id}"
                },
            },
        }

        if description:
            payload["description"] = {
                "format": "markdown",
                "raw": description,
            }

        if priority_id:
            payload["_links"]["priority"] = {
                "href": f"/api/v3/priorities/{priority_id}"
            }

        if estimated_hours is not None:
            payload["estimatedTime"] = f"PT{estimated_hours}H"

        if parent_id:
            payload["_links"]["parent"] = {
                "href": f"/api/v3/work_packages/{parent_id}"
            }

        response = self.session.post(
            f"{self.base_url}/api/v3/work_packages",
            json=payload,
        )
        response.raise_for_status()

        return response.json()


def load_config() -> dict:
    config_path = Path(__file__).with_name("config.yaml")

    with config_path.open("r", encoding="utf-8") as file:
        return yaml.safe_load(file)


def load_workbook_rows(workbook_path: Path) -> list[dict]:
    workbook = load_workbook(
        workbook_path,
        read_only=True,
        data_only=True,
    )

    if "OpenProject Import" not in workbook.sheetnames:
        raise ValueError(
            "Workbook is missing the 'OpenProject Import' worksheet."
        )

    worksheet = workbook["OpenProject Import"]
    rows = list(worksheet.iter_rows(values_only=True))

    if not rows:
        raise ValueError(
            "The 'OpenProject Import' worksheet is empty."
        )

    headers = set(rows[0])
    missing_columns = REQUIRED_COLUMNS - headers

    if missing_columns:
        raise ValueError(
            f"Workbook is missing required columns: "
            f"{sorted(missing_columns)}"
        )

    return [
        dict(zip(rows[0], row))
        for row in rows[1:]
        if any(value is not None for value in row)
    ]


def validate_rows(rows: list[dict]) -> None:
    valid_types = {"Feature", "Task"}

    features = {
        row["Subject"]
        for row in rows
        if row["Type"] == "Feature"
    }

    for row in rows:
        subject = row["Subject"]
        work_type = row["Type"]
        parent = row["Parent"]

        if work_type not in valid_types:
            raise ValueError(
                f"Invalid Type '{work_type}' for '{subject}'."
            )

        if work_type == "Feature" and parent:
            raise ValueError(
                f"Feature '{subject}' should not have a parent."
            )

        if work_type == "Task":
            if not parent:
                raise ValueError(
                    f"Task '{subject}' does not have a parent."
                )

            if parent not in features:
                raise ValueError(
                    f"Task '{subject}' references missing "
                    f"parent Feature '{parent}'."
                )


def find_existing_work_package(
    work_packages: list[dict],
    subject: str,
    work_type_name: str,
) -> dict | None:

    for work_package in work_packages:
        if work_package.get("subject") != subject:
            continue

        work_type = (
            work_package
            .get("_embedded", {})
            .get("type", {})
            .get("name")
        )

        if work_type == work_type_name:
            return work_package

    return None


def get_type_id(
    project: dict,
    work_type_name: str,
    client: OpenProjectClient,
) -> int:

    project_id = project["id"]

    response = client.session.get(
        f"{client.base_url}/api/v3/projects/{project_id}/types"
    )
    response.raise_for_status()

    types = response.json()["_embedded"]["elements"]

    for work_type in types:
        if work_type["name"].lower() == work_type_name.lower():
            return work_type["id"]

    raise ValueError(
        f"OpenProject type '{work_type_name}' is not enabled "
        f"for project '{project['identifier']}'."
    )


def get_priority_id(
    project: dict,
    priority_name: str,
    client: OpenProjectClient,
) -> int | None:

    response = client.session.get(
        f"{client.base_url}/api/v3/priorities"
    )
    response.raise_for_status()

    priorities = response.json()["_embedded"]["elements"]

    for priority in priorities:
        if priority["name"].lower() == priority_name.lower():
            return priority["id"]

    return None


def import_work_packages(
    rows: list[dict],
    client: OpenProjectClient,
    project: dict,
) -> None:

    project_id = project["id"]

    existing = client.get_work_packages(project_id)

    print()
    print("=" * 70)
    print("OPENPROJECT IMPORT")
    print("=" * 70)
    print(f"Project: {project['identifier']}")
    print(f"Existing work packages: {len(existing)}")
    print()

    features = [
        row for row in rows
        if row["Type"] == "Feature"
    ]

    tasks = [
        row for row in rows
        if row["Type"] == "Task"
    ]

    feature_ids = {}

    created_features = 0
    skipped_features = 0
    created_tasks = 0
    skipped_tasks = 0

    # ---------------------------------------------------------
    # Create Features
    # ---------------------------------------------------------

    print("FEATURES")
    print("-" * 70)

    for row in features:
        subject = row["Subject"]

        existing_feature = find_existing_work_package(
            existing,
            subject,
            "Feature",
        )

        if existing_feature:
            feature_ids[subject] = existing_feature["id"]
            skipped_features += 1

            print(f"[SKIP] {subject}")
            continue

        type_id = get_type_id(
            project,
            "Feature",
            client,
        )

        priority_id = get_priority_id(
            project,
            row["Priority"],
            client,
        )

        created = client.create_work_package(
            project_id=project_id,
            work_type_id=type_id,
            subject=subject,
            description=row["Description"],
            priority_id=priority_id,
            estimated_hours=row["Estimated hours"],
        )

        feature_ids[subject] = created["id"]
        created_features += 1

        print(f"[CREATE] {subject}")

    # Refresh work packages after creating Features.
    existing = client.get_work_packages(project_id)

    # ---------------------------------------------------------
    # Create Tasks
    # ---------------------------------------------------------

    print()
    print("TASKS")
    print("-" * 70)

    task_type_id = get_type_id(
        project,
        "Task",
        client,
    )

    for row in tasks:
        subject = row["Subject"]
        parent_subject = row["Parent"]

        parent_id = feature_ids.get(parent_subject)

        if not parent_id:
            raise ValueError(
                f"Could not resolve parent Feature "
                f"'{parent_subject}' for Task '{subject}'."
            )

        existing_task = find_existing_work_package(
            existing,
            subject,
            "Task",
        )

        if existing_task:
            skipped_tasks += 1

            print(f"[SKIP] {subject}")
            continue

        priority_id = get_priority_id(
            project,
            row["Priority"],
            client,
        )

        client.create_work_package(
            project_id=project_id,
            work_type_id=task_type_id,
            subject=subject,
            description=row["Description"],
            priority_id=priority_id,
            estimated_hours=row["Estimated hours"],
            parent_id=parent_id,
        )

        created_tasks += 1

        print(f"[CREATE] {subject}")

    print()
    print("=" * 70)
    print("IMPORT COMPLETE")
    print("=" * 70)
    print(f"Features created: {created_features}")
    print(f"Features skipped: {skipped_features}")
    print(f"Tasks created:    {created_tasks}")
    print(f"Tasks skipped:    {skipped_tasks}")
    print("=" * 70)


def main() -> None:
    config = load_config()

    api_token = os.environ.get("OPENPROJECT_API_TOKEN")

    if not api_token:
        raise RuntimeError(
            "OPENPROJECT_API_TOKEN environment variable is not set."
        )

    repo_root = Path(__file__).resolve().parents[2]

    workbook_path = repo_root / config["import"]["workbook"]

    if not workbook_path.exists():
        raise FileNotFoundError(
            f"Workbook not found: {workbook_path}"
        )

    rows = load_workbook_rows(workbook_path)

    validate_rows(rows)

    openproject_config = config["openproject"]

    client = OpenProjectClient(
        base_url=openproject_config["url"],
        api_token=api_token,
    )

    project = client.get_project(
        openproject_config["project_identifier"]
    )

    import_work_packages(
        rows=rows,
        client=client,
        project=project,
    )


if __name__ == "__main__":
    main()